from io import BytesIO
from typing import List, Literal

from fastapi import FastAPI
from fastapi.responses import Response
from pydantic import BaseModel, Field, field_validator
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

APP = FastAPI()
MAX_SLIDES = 40
MAX_BULLETS = 6
MAX_TITLE = 72
MAX_BULLET = 110

THEMES = {
    "corporate": {"bg": RGBColor(15, 32, 67), "accent": RGBColor(37, 99, 235), "title": RGBColor(255, 255, 255), "body": RGBColor(226, 232, 240)},
    "dark": {"bg": RGBColor(17, 24, 39), "accent": RGBColor(16, 185, 129), "title": RGBColor(255, 255, 255), "body": RGBColor(209, 213, 219)},
    "academic": {"bg": RGBColor(250, 250, 249), "accent": RGBColor(146, 64, 14), "title": RGBColor(28, 25, 23), "body": RGBColor(68, 64, 60)},
    "creative": {"bg": RGBColor(76, 29, 149), "accent": RGBColor(251, 191, 36), "title": RGBColor(255, 255, 255), "body": RGBColor(237, 233, 254)},
}
FONTS = {"ar": "Arial", "en": "Calibri", "fr": "Calibri"}


class SlideSpec(BaseModel):
    title: str = ""
    bullets: List[str] = Field(default_factory=list)


class PresentationSpec(BaseModel):
    filename: str = "presentation.pptx"
    title: str = ""
    language: Literal["ar", "en", "fr"] = "ar"
    theme: Literal["corporate", "dark", "academic", "creative"] = "corporate"
    slide_count: int = 8
    slides: List[SlideSpec]

    @field_validator("slide_count")
    @classmethod
    def clamp_count(cls, v: int) -> int:
        return max(1, min(int(v), MAX_SLIDES))


def clip(text: str, n: int) -> str:
    t = " ".join((text or "").split())
    return t if len(t) <= n else t[: n - 1].rstrip() + "…"


def set_rtl(paragraph, enabled: bool) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    for child in list(pPr):
        if child.tag == qn("a:rtl"):
            pPr.remove(child)
    rtl = etree.SubElement(pPr, qn("a:rtl"))
    rtl.set("val", "1" if enabled else "0")


def fill_bg(slide, color: RGBColor) -> None:
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.line.fill.background()
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    spTree = slide.shapes._spTree
    sp = bg._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def write_box(tf, lines, font_name, size, color, bold, align, rtl):
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.clear()
        run = p.add_run()
        run.text = line
        run.font.name = font_name
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        p.alignment = align
        p.space_after = Pt(8)
        p.level = 0
        set_rtl(p, rtl)


@APP.post("/generate")
def generate(spec: PresentationSpec):
    slides = spec.slides[: spec.slide_count]
    while len(slides) < spec.slide_count:
        slides.append(SlideSpec(title=spec.title or "Slide", bullets=["—"]))

    theme = THEMES[spec.theme]
    font = FONTS[spec.language]
    rtl = spec.language == "ar"
    align = PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for idx, s in enumerate(slides):
        slide = prs.slides.add_slide(blank)
        fill_bg(slide, theme["bg"])

        accent = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(7.5)
        )
        accent.line.fill.background()
        accent.fill.solid()
        accent.fill.fore_color.rgb = theme["accent"]

        title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(12.1), Inches(1.15))
        write_box(title_box.text_frame, [clip(s.title, MAX_TITLE)], font, 28, theme["title"], True, align, rtl)

        bullets = [clip(b, MAX_BULLET) for b in (s.bullets or []) if str(b).strip()][:MAX_BULLETS] or ["—"]
        body = slide.shapes.add_textbox(Inches(0.85), Inches(1.7), Inches(11.8), Inches(5.1))
        write_box(body.text_frame, ["• " + b for b in bullets], font, 20, theme["body"], False, align, rtl)

        foot = slide.shapes.add_textbox(Inches(0.7), Inches(7.05), Inches(12.1), Inches(0.3))
        write_box(
            foot.text_frame,
            [f"{idx + 1} / {len(slides)}"],
            font,
            11,
            theme["body"],
            False,
            PP_ALIGN.LEFT if rtl else PP_ALIGN.RIGHT,
            False,
        )

    buf = BytesIO()
    prs.save(buf)
    name = spec.filename if spec.filename.endswith(".pptx") else spec.filename + ".pptx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )
